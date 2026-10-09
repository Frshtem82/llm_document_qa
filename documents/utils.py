"""توابع کمکی برای پردازش اسناد متنی و فایل‌های docx."""
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter


def extract_text_from_docx(file_path: str) -> str:
    doc = DocxDocument(file_path)
    parts = []
    
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    parts.append(text)
    
    return "\n".join(parts)


def extract_text_from_file(file_field) -> str:
    if not file_field:
        return ""
    
    name = file_field.name.lower()
    file_path = file_field.path
    
    if name.endswith('.docx'):
        return extract_text_from_docx(file_path)
    elif name.endswith('.txt'):
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    else:
        raise ValueError(f"فرمت فایل پشتیبانی نمی‌شه: {name}")


def split_text_into_chunks(text: str, chunk_size: int = 500, chunk_overlap: int = 50) -> list:
    if not text or not text.strip():
        return []
    
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        length_function=len,
        separators=["\n\n", "\n", ".", "،", " ", ""],
    )
    
    chunks = splitter.split_text(text)
    return [c.strip() for c in chunks if c.strip()]